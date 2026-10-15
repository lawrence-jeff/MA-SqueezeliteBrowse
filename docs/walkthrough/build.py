from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SHOTS = Path(__file__).parent / "shots"
# (file, title, note)
slides = [
    ("01", "Main Menu", "Home screen of the piCorePlayer (JiveLite). Settings holds the HA Extras entry when it is enabled for the device."),
    ("23", "Now Playing", "Reached with the Now Playing tile or the [ key. Shows the current track from the Music Assistant queue."),
    ("24", "Queue (Current Playlist)", "The player's queue as Music Assistant has it. Long press a row for Play Now, Play Next, Move to End and Delete item."),
    ("02", "My Music", "The two preset tiles come from the player's settings in Music Assistant (favorited playlists and radio stations)."),
    ("03", "My Music - Favorites", "All Favorites, then one entry per type."),
    ("04", "My Music - Artists", ""),
    ("22", "My Music - Artists - Avril Lavigne", "An artist who is only on one track (a featured artist). The view lists albums, then tracks not on those albums, then matching radio stations."),
    ("05", "My Music - Albums", ""),
    ("06", "My Music - Albums - an album", "Tracks of one album. Tapping a track adds it (plays it when the queue is idle); long press for Play Album from here."),
    ("07", "My Music - Tracks", ""),
    ("08", "My Music - Playlists", ""),
    ("09", "My Music - Audiobooks", ""),
    ("10", "My Music - Podcasts", ""),
    ("11", "My Music - Podcasts - a podcast", "Episodes of one podcast."),
    ("12", "My Music - Radio", ""),
    ("13", "My Music - Search", "The search text entry."),
    ("14", "My Music - Switch Library", ""),
    ("15", "Settings", "Includes Alarm Clock (says alarms are not supported) and HA Extras."),
    ("16", "Settings - HA Extras", "Experimental. Needs the device setting 'Expose HA Entities to Device (HA Extras)' and the HA Extras label on the entities."),
    ("17", "Settings - HA Extras - Cameras", "Choosing a camera shows a snapshot full screen in the Image Viewer, refreshed about every 10 seconds."),
    ("27", "Settings - HA Extras - Cameras - Doorbell", "The camera view in the Image Viewer. It fetches a new snapshot about every 10 seconds until a key is pressed."),
    ("18", "Settings - HA Extras - Dashboards", "Entities whose state is the address of a picture of a dashboard. Shown full screen like the cameras."),
    ("28", "Settings - HA Extras - Dashboards - Squeezebox", "A dashboard in the Image Viewer. A screenshot service renders the picture; it is fetched again about every 10 seconds."),
    ("19", "Settings - HA Extras - Controls", "Choosing a row opens Turn On / Turn Off / Cancel."),
    ("25", "Settings - HA Extras - Controls - Porch Lights", "Choosing a control opens its menu: the current state, Turn On, Turn Off and Cancel."),
    ("20", "Settings - HA Extras - Scenes and Scripts", "Choosing one opens Run / Cancel."),
    ("26", "Settings - HA Extras - Scenes and Scripts - Test Scene", "Choosing a scene opens Run and Cancel."),
    ("21", "Settings - HA Extras - Status", "Read-only values."),
]
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

def text(slide, left, top, width, height, value, size, bold=False, color=(0x22, 0x22, 0x22)):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = value
    p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = RGBColor(*color)
    return box

s = prs.slides.add_slide(blank)
text(s, 0.8, 2.2, 11.7, 1.2, "MA-SqueezeliteBrowse: menu walkthrough", 40, True)
text(s, 0.8, 3.5, 11.7, 2.5,
     "Screens of the piCorePlayer (JiveLite, 1920x1080) browsing Music Assistant 2.10.5 with the patched Squeezelite provider.\n"
     "Captured 2026-10-11. The content shown is a moment in time: it follows the library, the presets and the Home Assistant labels.", 18)

for key, title, note in slides:
    s = prs.slides.add_slide(blank)
    text(s, 0.5, 0.2, 12.3, 0.7, title, 28, True)
    height = 5.9
    width = height * 16 / 9
    s.shapes.add_picture(str(SHOTS / f"{key}.jpg"), Inches((13.333 - width) / 2), Inches(1.0), Inches(width), Inches(height))
    if note:
        text(s, 0.5, 6.95, 12.3, 0.5, note, 12, color=(0x55, 0x55, 0x55))
out = Path(__file__).parent / "MA-SqueezeliteBrowse-menu-walkthrough.pptx"
prs.save(out)
print(out, out.stat().st_size // 1024, "KB", len(prs.slides), "slides")
